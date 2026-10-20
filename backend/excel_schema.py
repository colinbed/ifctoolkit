"""Executable Excel choices come exclusively from the installed EXPRESS schema."""
from functools import lru_cache

import ifcopenshell
import ifcopenshell.ifcopenshell_wrapper as wrapper
import ifcopenshell.util.type
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

REFERENCE_VERSION = f"2.0 / IfcOpenShell {ifcopenshell.version}"


def inherits(declaration, name):
    while declaration:
        if declaration.name() == name:
            return True
        declaration = declaration.supertype()
    return False


@lru_cache(maxsize=8)
def schema_reference(schema):
    definition = wrapper.schema_by_name(schema)
    result = {}
    for decl in definition.entities():
        category = "type" if inherits(decl, "IfcTypeObject") else "occurrence" if inherits(decl, "IfcProduct") else "reference"
        attributes = {a.name(): a for a in decl.all_attributes()}
        enum = []
        if "PredefinedType" in attributes:
            enum = list(attributes["PredefinedType"].type_of_attribute().declared_type().enumeration_items())
        correspondence = []
        if category == "occurrence":
            correspondence = ifcopenshell.util.type.get_applicable_types(decl.name(), schema) or []
        elif category == "type":
            correspondence = ifcopenshell.util.type.get_applicable_entities(decl.name(), schema) or []
        # reassign_class supports element/type conversions with a known counterpart.
        required = tuple(a.name() for a in decl.all_attributes() if not a.optional())
        # Required construction/operation/etc attributes cannot be supplied by
        # this workbook, so those classes are reference-only conversion targets.
        supported_required = {"GlobalId", "OwnerHistory", "PredefinedType"}
        selectable = not decl.is_abstract() and bool(correspondence) and not (set(required) - supported_required) and (inherits(decl, "IfcElement") or inherits(decl, "IfcElementType") or decl.name() in {"IfcDoorStyle", "IfcWindowStyle"})
        result[decl.name()] = {"entity": decl.name(), "category": category,
            "abstract": decl.is_abstract(), "selectable": selectable,
            "attributes": tuple(attributes), "required": required, "enum": tuple(enum), "correspondence": tuple(correspondence)}
    return result


def documentation_url(schema, entity):
    if schema.upper() == "IFC2X3":
        return f"https://standards.buildingsmart.org/IFC/RELEASE/IFC2x3/TC1/HTML/"
    release = "IFC4x3_ADD2" if schema.upper().startswith("IFC4X3") else "IFC4_ADD2_TC1"
    return f"https://ifc43-docs.standards.buildingsmart.org/{release}/HTML/lexical/{entity}.htm" if schema.upper().startswith("IFC4X3") else f"https://standards.buildingsmart.org/IFC/RELEASE/IFC4/ADD2_TC1/HTML/link/{entity.lower()}.htm"


def add_schema_sheets(workbook, sheet, schema, row_counts, columns, row_entities=None):
    reference = schema_reference(schema)
    ws = sheet("SchemaEntities", ["Schema", "IFC Entity", "Category", "Abstract", "Selectable", "Correspondence", "Native PredefinedType", "PredefinedType Values", "Documentation URL", "Reference Version", "Description"])
    # Descriptions are read from IfcOpenShell's bundled buildingSMART documentation.
    import ifcopenshell.util.doc
    for name, info in sorted(reference.items()):
        try:
            doc = ifcopenshell.util.doc.get_entity_doc(schema, name) or {}
        except (KeyError, RuntimeError):
            doc = {}
        ws.append([schema, name, info["category"], info["abstract"], info["selectable"], ", ".join(info["correspondence"]), "PredefinedType" in info["attributes"], ", ".join(info["enum"]), doc.get("spec_url") or documentation_url(schema, name), REFERENCE_VERSION, (doc.get("description") or "See official documentation link.")[:32000]])
    lists = sheet("SchemaChoices", ["Range", "Value"])
    lists.append(["NoPredefinedType", None])
    workbook.defined_names.add(DefinedName("PT_Unavailable", attr_text="'SchemaChoices'!$B$2:$B$2"))
    cursor = 3
    for category, range_name in (("occurrence", "OccurrenceEntities"), ("type", "TypeEntities")):
        names = sorted(n for n, i in reference.items() if i["category"] == category and i["selectable"])
        start = cursor
        for name in names:
            lists.append([range_name, name]); cursor += 1
        workbook.defined_names.add(DefinedName(range_name, attr_text=f"'SchemaChoices'!$B${start}:$B${max(start, cursor - 1)}"))
    for name, info in sorted(reference.items()):
        start = cursor
        for value in info["enum"] or (None,):
            lists.append([name, value]); cursor += 1
        workbook.defined_names.add(DefinedName(f"PT_{name}", attr_text=f"'SchemaChoices'!$B${start}:$B${cursor - 1}"))
    for title, count in row_counts.items():
        if not count or title not in columns:
            continue
        ws = workbook[title]
        headers = columns[title]
        entity_col = get_column_letter(headers.index("IFC Entity") + 1)
        predef_col = get_column_letter(headers.index("Predefined Type") + 1)
        dv = DataValidation(type="list", formula1="=TypeEntities" if title == "Types" else "=OccurrenceEntities", allow_blank=True)
        dv.errorTitle = "Invalid IFC class"; dv.error = "Choose a concrete class from this schema."; dv.showErrorMessage = True; dv.errorStyle = "stop"
        objects = (row_entities or {}).get(title)
        if objects is None:
            dv.add(f"{entity_col}2:{entity_col}{count + 1}")
        else:
            start = None
            for row, obj in enumerate(objects, 2):
                eligible = reference[obj.is_a()]["selectable"]
                if eligible and start is None:
                    start = row
                if start is not None and (not eligible or row == count + 1):
                    end = row if eligible else row - 1
                    dv.add(f"{entity_col}{start}:{entity_col}{end}")
                    start = None
        if dv.sqref:
            ws.data_validations.append(dv)
        dv = DataValidation(type="list", formula1=f'INDIRECT("PT_"&${entity_col}2)', allow_blank=True)
        dv.errorTitle = "Invalid PredefinedType"; dv.error = "Choose a value supported by this IFC class; blank means unavailable/no change."; dv.showErrorMessage = True; dv.errorStyle = "stop"
        dv.add(f"{predef_col}2:{predef_col}{count + 1}"); ws.data_validations.append(dv)
