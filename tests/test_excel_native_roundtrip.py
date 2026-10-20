"""End-to-end assertions reopen both XLSX and IFC, including native type links."""
import hashlib

import ifcopenshell
import ifcopenshell.util.element as util
import pytest
from ifcopenshell.guid import new as guid
from openpyxl import load_workbook

from app import extract_to_excel, update_ifc_from_excel, ExcelWritebackValidationError
from backend.excel_schema import schema_reference
from ifc_protocol import default_protocol_config


def model(schema="IFC4"):
    f = ifcopenshell.file(schema=schema)
    project = f.create_entity("IfcProject", GlobalId=guid(), Name="Project", LongName="P1")
    building = f.create_entity("IfcBuilding", GlobalId=guid(), Name="Building")
    storey = f.create_entity("IfcBuildingStorey", GlobalId=guid(), Name="Floor", Description="Floor desc", LongName="Level 1", Elevation=12.)
    f.create_entity("IfcRelAggregates", GlobalId=guid(), RelatingObject=project, RelatedObjects=[building])
    f.create_entity("IfcRelAggregates", GlobalId=guid(), RelatingObject=building, RelatedObjects=[storey])
    typ = f.create_entity("IfcBeamType", GlobalId=guid(), Name="Beam type", Description="Type desc", PredefinedType="BEAM")
    a = f.create_entity("IfcBeam", GlobalId=guid(), Name="Beam A", Description="A desc", ObjectType="Beam")
    b = f.create_entity("IfcBeam", GlobalId=guid(), Name="Beam B", Description="B desc", ObjectType="Beam")
    if hasattr(a, "PredefinedType"):
        a.PredefinedType = b.PredefinedType = "BEAM"
    f.create_entity("IfcRelDefinesByType", GlobalId=guid(), RelatedObjects=[a, b], RelatingType=typ)
    f.create_entity("IfcRelContainedInSpatialStructure", GlobalId=guid(), RelatingStructure=storey, RelatedElements=[a, b])
    point = f.create_entity("IfcCartesianPoint", Coordinates=(0., 0., 0.))
    axis = f.create_entity("IfcAxis2Placement3D", Location=point)
    a.ObjectPlacement = f.create_entity("IfcLocalPlacement", RelativePlacement=axis)
    context = f.create_entity("IfcGeometricRepresentationContext", CoordinateSpaceDimension=3, WorldCoordinateSystem=axis)
    line = f.create_entity("IfcPolyline", Points=[point, f.create_entity("IfcCartesianPoint", Coordinates=(1., 0., 0.))])
    rep = f.create_entity("IfcShapeRepresentation", ContextOfItems=context, RepresentationIdentifier="Body", RepresentationType="Curve3D", Items=[line])
    a.Representation = f.create_entity("IfcProductDefinitionShape", Representations=[rep])
    return f, a, b, typ, storey


def export(tmp_path, schema="IFC4", protocol=False, prepare=None):
    f, a, b, typ, storey = model(schema)
    if prepare:
        prepare(f, a, b, typ, storey)
    source, xlsx, output = [tmp_path / n for n in ("source.ifc", "edit.xlsx", "edited.ifc")]
    f.write(str(source))
    extract_to_excel(str(source), str(xlsx), {"protocol": default_protocol_config()} if protocol else None)
    return source, xlsx, output, (a.GlobalId, b.GlobalId, typ.GlobalId, storey.GlobalId)


def edit(xlsx, sheet, identity, column, value):
    wb = load_workbook(xlsx)
    ws = wb[sheet]
    headers = {c.value: c.column for c in ws[1]}
    idcol = headers["Type GlobalId" if sheet == "Types" else "GlobalId"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, idcol).value == identity)
    if column not in headers:
        headers[column] = ws.max_column + 1
        ws.cell(1, headers[column], column)
    ws.cell(row, headers[column], value)
    wb.save(xlsx); wb.close()


def apply(paths, **kwargs):
    source, xlsx, output, ids = paths
    digest = hashlib.sha256(source.read_bytes()).digest()
    result = update_ifc_from_excel(str(source), str(xlsx), str(output), **kwargs)
    assert hashlib.sha256(source.read_bytes()).digest() == digest
    return ifcopenshell.open(str(output)), result


def test_storey_native_properties_and_distinct_classification(tmp_path):
    paths = export(tmp_path)
    source, xlsx, output, (a, b, typ, floor) = paths
    for column, value in {"Name": "Edited floor", "Description": "Edited description", "Long Name": "Edited level", "COBie_Floor.Category": "Floor category", "Classification: COBie Floors Reference": "FF-01", "Classification: COBie Floors Name": "Floor classification"}.items():
        edit(xlsx, "BuildingStoreys", floor, column, value)
    f, result = apply(paths)
    storey = f.by_guid(floor)
    assert (storey.Name, storey.Description, storey.LongName, storey.Elevation) == ("Edited floor", "Edited description", "Edited level", 12.)
    assert util.get_psets(storey)["COBie_Floor"]["Category"] == "Floor category"
    assert storey.HasAssociations[0].RelatingClassification.Identification == "FF-01"
    assert f.by_guid(a).Name == "Beam A"
    assert util.get_container(f.by_guid(a)).GlobalId == floor
    assert result["applied"] >= 6


@pytest.mark.parametrize("schema", ["IFC2X3", "IFC4", "IFC4X3_ADD2"])
def test_instance_and_assigned_type_targets_and_protocol_does_not_revert(tmp_path, schema):
    paths = export(tmp_path, schema, protocol=True)
    _, xlsx, _, (a, b, typ, _) = paths
    for col, value in {"Instance Name": "Changed A", "Instance Description": "Changed A desc", "Type Name": "Changed type", "Type Description": "Changed type desc"}.items():
        edit(xlsx, "Elements", a, col, value)
    f, result = apply(paths)
    assert (f.by_guid(a).Name, f.by_guid(a).Description) == ("Changed A", "Changed A desc")
    assert (f.by_guid(b).Name, f.by_guid(b).Description) == ("Beam B", "B desc")
    assigned = util.get_type(f.by_guid(a))
    assert assigned.GlobalId == typ
    assert (assigned.Name, assigned.Description) == ("Changed type", "Changed type desc")
    assert f.schema_identifier == schema


def test_overlapping_type_edits_consistent_and_conflicting(tmp_path):
    paths = export(tmp_path)
    _, xlsx, output, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "Type Name", "New type")
    edit(xlsx, "Elements", b, "Type Name", "New type")
    edit(xlsx, "Types", typ, "Type Name", "New type")
    f, _ = apply(paths)
    assert f.by_guid(typ).Name == "New type"
    output.unlink()
    edit(xlsx, "Types", typ, "Type Name", "Conflict")
    with pytest.raises(ExcelWritebackValidationError, match="Conflicting.*edits"):
        apply(paths)
    assert not output.exists()


def test_authorised_shared_type_split(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "Type Name", "A type")
    edit(xlsx, "Elements", b, "Type Name", "B type")
    with pytest.raises(ExcelWritebackValidationError):
        apply(paths)
    f, _ = apply(paths, split_shared_types=True)
    assert util.get_type(f.by_guid(a)).Name == "A type"
    assert util.get_type(f.by_guid(b)).Name == "B type"
    assert util.get_type(f.by_guid(a)) != util.get_type(f.by_guid(b))


def test_simple_shared_material_reassociation(tmp_path):
    def prepare(f, a, b, typ, storey):
        mat = f.create_entity("IfcMaterial", Name="Steel")
        f.create_entity("IfcRelAssociatesMaterial", GlobalId=guid(), RelatedObjects=[a, b], RelatingMaterial=mat)
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, _, (a, b, _, _) = paths
    edit(xlsx, "Elements", a, "Material", "Concrete")
    f, _ = apply(paths)
    assert util.get_material(f.by_guid(a)).Name == "Concrete"
    assert util.get_material(f.by_guid(b)).Name == "Steel"
    assert len(f.by_type("IfcMaterial")) == 2


def test_type_scoped_material_conflicts(tmp_path):
    def prepare(f, a, b, typ, storey):
        mat = f.create_entity("IfcMaterial", Name="Steel")
        f.create_entity("IfcRelAssociatesMaterial", GlobalId=guid(), RelatedObjects=[typ], RelatingMaterial=mat)
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, _, (a, b, _, _) = paths
    edit(xlsx, "Elements", a, "Material", "Concrete")
    edit(xlsx, "Elements", b, "Material", "Timber")
    with pytest.raises(ExcelWritebackValidationError, match="Conflicting.*edits"):
        apply(paths)


@pytest.mark.parametrize("kind", ["IfcMaterialLayerSet", "IfcMaterialList", "IfcMaterialProfileSet", "IfcMaterialConstituentSet"])
def test_structured_material_reference_only(tmp_path, kind):
    def prepare(f, a, b, typ, storey):
        mat = f.create_entity("IfcMaterial", Name="Steel")
        if kind == "IfcMaterialLayerSet":
            layer = f.create_entity("IfcMaterialLayer", Material=mat, LayerThickness=.2)
            structure = f.create_entity(kind, MaterialLayers=[layer], LayerSetName="Layers")
        elif kind == "IfcMaterialList":
            structure = f.create_entity(kind, Materials=[mat])
        elif kind == "IfcMaterialProfileSet":
            structure = f.create_entity(kind, Name="Profiles", MaterialProfiles=[])
        else:
            structure = f.create_entity(kind, Name="Constituents", MaterialConstituents=[])
        f.create_entity("IfcRelAssociatesMaterial", GlobalId=guid(), RelatedObjects=[a], RelatingMaterial=structure)
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, output, (a, *_) = paths
    f, _ = apply(paths)
    assert util.get_material(f.by_guid(a)).is_a() == kind
    output.unlink()
    edit(xlsx, "Elements", a, "Material", "Flattened")
    with pytest.raises(ExcelWritebackValidationError, match="Structured material"):
        apply(paths)
    assert not output.exists()


@pytest.mark.parametrize("schema", ["IFC2X3", "IFC4", "IFC4X3_ADD2"])
def test_exact_schema_references_and_dependent_dropdowns(tmp_path, schema):
    paths = export(tmp_path, schema)
    _, xlsx, _, ids = paths
    wb = load_workbook(xlsx)
    headers = [c.value for c in wb["Elements"][1]]
    assert headers[:9] == ["GlobalId", "Instance Name", "Instance Description", "Type Name", "Type Description", "Object Type", "Material", "IFC Entity", "Predefined Type"]
    assert "OccurrenceName" not in headers and "TypeDescription" not in headers
    ref = schema_reference(schema)
    occurrence_range = wb.defined_names["OccurrenceEntities"]
    _, coord = next(occurrence_range.destinations)
    choices = [r[0].value for r in wb["SchemaChoices"][coord]]
    assert "IfcWall" in choices and "IfcElement" not in choices and "IfcBeamType" not in choices
    assert "IfcWall" not in {row[0].value for row in wb["SchemaChoices"][next(wb.defined_names["TypeEntities"].destinations)[1]]}
    for title, name in (("Elements", "=OccurrenceEntities"), ("Types", "=TypeEntities")):
        dvs = list(wb[title].data_validations.dataValidation)
        assert any(d.formula1 == name for d in dvs)
        dependent = next(d for d in dvs if "INDIRECT" in d.formula1)
        assert '"PT_"' in dependent.formula1
        assert dependent.sqref
    _, enum_range = next(wb.defined_names["PT_IfcBeamType"].destinations)
    assert tuple(r[0].value for r in wb["SchemaChoices"][enum_range]) == ref["IfcBeamType"]["enum"]
    assert wb["SchemaEntities"].sheet_state == "visible"
    wb.save(xlsx); wb.close()
    apply(paths)


def test_ifc2x3_type_predefined_and_occurrence_unavailable(tmp_path):
    paths = export(tmp_path, "IFC2X3")
    _, xlsx, output, (a, b, typ, _) = paths
    edit(xlsx, "Types", typ, "Predefined Type", "JOIST")
    f, _ = apply(paths)
    assert util.get_type(f.by_guid(a)).PredefinedType == "JOIST"
    assert not hasattr(f.by_guid(a), "PredefinedType")
    output.unlink()
    edit(xlsx, "Elements", a, "Predefined Type", "JOIST")
    with pytest.raises(ExcelWritebackValidationError, match="no native PredefinedType"):
        apply(paths)


@pytest.mark.parametrize("value", ["MADE_UP", "USERDEFINED"])
def test_invalid_pasted_enum_and_missing_companion_rejected(tmp_path, value):
    paths = export(tmp_path)
    _, xlsx, output, (_, _, typ, _) = paths
    edit(xlsx, "Types", typ, "Predefined Type", value)
    with pytest.raises(ExcelWritebackValidationError):
        apply(paths)
    assert not output.exists()


def test_userdefined_companion(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (_, _, typ, _) = paths
    edit(xlsx, "Types", typ, "Predefined Type", "USERDEFINED")
    edit(xlsx, "Types", typ, "Element Type", "Special beam")
    f, _ = apply(paths)
    assert f.by_guid(typ).PredefinedType == "USERDEFINED"
    assert f.by_guid(typ).ElementType == "Special beam"


def test_type_class_conversion_preserves_guid_geometry_and_relationships(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (a, b, typ, storey) = paths
    source = ifcopenshell.open(str(paths[0]))
    representation = str(source.by_guid(a).Representation)
    placement = str(source.by_guid(a).ObjectPlacement)
    edit(xlsx, "Types", typ, "IFC Entity", "IfcColumnType")
    edit(xlsx, "Types", typ, "Predefined Type", "COLUMN")
    for occurrence in (a, b):
        edit(xlsx, "Elements", occurrence, "Predefined Type", "COLUMN")
    f, _ = apply(paths)
    assert f.by_guid(typ).is_a() == "IfcColumnType"
    for identity in (a, b):
        assert f.by_guid(identity).is_a() == "IfcColumn"
        assert util.get_type(f.by_guid(identity)).GlobalId == typ
        assert util.get_container(f.by_guid(identity)).GlobalId == storey
    assert str(f.by_guid(a).Representation) == representation
    assert str(f.by_guid(a).ObjectPlacement) == placement


def test_unsupported_conversion_and_incompatible_assigned_type_rejected(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (a, _, _, _) = paths
    for target in ("IfcElement", "IfcProject", "IfcColumn", "IfcDoesNotExist"):
        edit(xlsx, "Elements", a, "IFC Entity", target)
        with pytest.raises(ExcelWritebackValidationError):
            apply(paths)


def test_version1_aliases_and_conflicting_aliases(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (a, _, typ, _) = paths
    wb = load_workbook(xlsx)
    ws = wb["Elements"]
    mapping = {"Instance Name": "OccurrenceName", "Instance Description": "TypeDescription", "Type Name": "TypeName"}
    for cell in ws[1]:
        if cell.value in mapping: cell.value = mapping[cell.value]
    wb.save(xlsx); wb.close()
    edit(xlsx, "Elements", a, "OccurrenceName", "Legacy name")
    edit(xlsx, "Elements", a, "TypeDescription", "Legacy occurrence description")
    edit(xlsx, "Elements", a, "TypeName", "Legacy type name")
    f, _ = apply(paths)
    assert f.by_guid(a).Description == "Legacy occurrence description"
    assert f.by_guid(typ).Description == "Type desc"
    assert f.by_guid(typ).Name == "Legacy type name"
    edit(xlsx, "Elements", a, "Instance Name", "Conflicting new name")
    with pytest.raises(ExcelWritebackValidationError, match="Conflicting.*edits"):
        apply(paths)


def test_protocol_explicit_conflict_and_protocol_only_edit(tmp_path):
    paths = export(tmp_path, protocol=True)
    _, xlsx, _, (a, _, _, _) = paths
    edit(xlsx, "ProtocolData", a, "Asset Name", "Protocol name")
    f, _ = apply(paths)
    assert f.by_guid(a).Name == "Protocol name"
    edit(xlsx, "Elements", a, "Instance Name", "Other name")
    with pytest.raises(ExcelWritebackValidationError, match="Conflicting.*edits"):
        apply(paths)


def test_existing_floor_properties_export_and_shared_pset_isolated(tmp_path):
    def prepare(f, a, b, typ, storey):
        prop = f.create_entity("IfcPropertySingleValue", Name="Category", NominalValue=f.create_entity("IfcLabel", "Existing category"))
        pset = f.create_entity("IfcPropertySet", GlobalId=guid(), Name="COBie_Floor", HasProperties=[prop])
        f.create_entity("IfcRelDefinesByProperties", GlobalId=guid(), RelatedObjects=[storey, a], RelatingPropertyDefinition=pset)
        # A numeric property must retain its IFC measure type.
        prop2 = f.create_entity("IfcPropertySingleValue", Name="HeightToHaunch", NominalValue=f.create_entity("IfcLengthMeasure", 3.))
        pset2 = f.create_entity("IfcPropertySet", GlobalId=guid(), Name="COBie_FloorAttributes", HasProperties=[prop2])
        f.create_entity("IfcRelDefinesByProperties", GlobalId=guid(), RelatedObjects=[storey], RelatingPropertyDefinition=pset2)
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, _, (a, b, typ, floor) = paths
    wb = load_workbook(xlsx)
    ws = wb["BuildingStoreys"]
    headers = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, headers["COBie_Floor.Category"]).value == "Existing category"
    wb.close()
    edit(xlsx, "BuildingStoreys", floor, "COBie_Floor.Category", "Changed category")
    edit(xlsx, "BuildingStoreys", floor, "COBie_FloorAttributes.HeightToHaunch", 4.)
    f, _ = apply(paths)
    assert util.get_psets(f.by_guid(floor))["COBie_Floor"]["Category"] == "Changed category"
    assert util.get_psets(f.by_guid(a))["COBie_Floor"]["Category"] == "Existing category"
    pset = util.get_psets(f.by_guid(floor))["COBie_FloorAttributes"]
    prop = f.by_id(pset["id"]).HasProperties[0]
    assert prop.NominalValue.is_a() == "IfcLengthMeasure"
    assert prop.NominalValue.wrappedValue == 4.


def test_single_material_on_type_is_applied_at_type_and_preserves_other_assignments(tmp_path):
    def prepare(f, a, b, typ, storey):
        mat = f.create_entity("IfcMaterial", Name="Steel")
        unrelated = f.create_entity("IfcBeam", GlobalId=guid(), Name="Unrelated")
        f.create_entity("IfcRelAssociatesMaterial", GlobalId=guid(), RelatedObjects=[typ, unrelated], RelatingMaterial=mat)
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, _, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "Material", "Timber")
    f, _ = apply(paths)
    assert util.get_material(f.by_guid(typ), should_inherit=False).Name == "Timber"
    assert util.get_material(f.by_guid(b)).Name == "Timber"
    assert util.get_material(f.by_guid(a), should_inherit=False) is None
    assert util.get_material(next(o for o in f.by_type("IfcBeam") if o.Name == "Unrelated")).Name == "Steel"


def test_occurrence_only_conversion_preserves_geometry(tmp_path):
    def prepare(f, a, b, typ, storey):
        # An untyped occurrence can change class independently.
        rel = next(r for r in f.by_type("IfcRelDefinesByType") if a in r.RelatedObjects)
        rel.RelatedObjects = [b]
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, _, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "IFC Entity", "IfcColumn")
    edit(xlsx, "Elements", a, "Predefined Type", "COLUMN")
    f, _ = apply(paths)
    assert f.by_guid(a).is_a() == "IfcColumn"
    assert f.by_guid(b).is_a() == "IfcBeam"
    assert f.by_guid(typ).is_a() == "IfcBeamType"
    assert f.by_guid(a).Representation is not None


def test_reordered_rows_resolve_by_globalid(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "Instance Name", "Edited A")
    wb = load_workbook(xlsx)
    ws = wb["Elements"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    ws.delete_rows(2, ws.max_row)
    for row in reversed(rows): ws.append(row)
    wb.save(xlsx); wb.close()
    f, _ = apply(paths)
    assert f.by_guid(a).Name == "Edited A"
    assert f.by_guid(b).Name == "Beam B"


def test_verification_failure_keeps_existing_output_and_source(tmp_path, monkeypatch):
    from backend.excel_roundtrip import RoundTripPlan
    paths = export(tmp_path)
    source, xlsx, output, (a, *_) = paths
    output.write_text("Existing output")
    edit(xlsx, "Elements", a, "Instance Name", "Edited")
    before = source.read_bytes()
    def fail(*_):
        raise ValueError("Verification intentionally failed")
    monkeypatch.setattr(RoundTripPlan, "verify", fail)
    with pytest.raises(ValueError, match="Verification intentionally failed"):
        apply(paths)
    assert output.read_text() == "Existing output"
    assert source.read_bytes() == before
    assert not list(tmp_path.glob(".ifc-write-*"))


def test_legacy_authorised_mapping_still_uses_batch_planner(tmp_path):
    def prepare(f, a, b, typ, storey):
        f.by_type("IfcRelDefinesByType")[0].RelatedObjects = [b]
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, _, (a, *_) = paths
    edit(xlsx, "Elements", a, "ApplyChange", "Yes")
    edit(xlsx, "Elements", a, "TargetPredefinedType", "COLUMN")
    calls = []
    def infer(signature, row):
        calls.append(signature)
        return "IfcColumn"
    f, _ = apply(paths, proxy_mapping_inference=infer)
    assert len(calls) == 1
    assert f.by_guid(a).is_a() == "IfcColumn"
    assert f.by_guid(a).PredefinedType == "COLUMN"


def test_storey_class_conversion_rejected_before_mutation(tmp_path):
    paths = export(tmp_path)
    _, xlsx, output, (_, _, _, floor) = paths
    edit(xlsx, "Elements", floor, "IFC Entity", "IfcWall")
    with pytest.raises(ExcelWritebackValidationError, match="Unsupported conversion"):
        apply(paths)
    assert not output.exists()


def test_add_new_does_not_create_blank_site_or_reparent_existing_building(tmp_path):
    paths = export(tmp_path)
    _, xlsx, _, (_, _, _, floor) = paths
    edit(xlsx, "BuildingStoreys", floor, "Name", "Edited floor")
    f, _ = apply(paths, add_new="yes")
    assert not f.by_type("IfcSite")
    building = f.by_type("IfcBuilding")[0]
    assert util.get_aggregate(building).is_a() == "IfcProject"


def test_repeated_type_material_import_does_not_restore_unchanged_aliases(tmp_path):
    def prepare(f, a, b, typ, storey):
        mat = f.create_entity("IfcMaterial", Name="Steel")
        f.create_entity("IfcRelAssociatesMaterial", GlobalId=guid(), RelatedObjects=[typ], RelatingMaterial=mat)
    paths = export(tmp_path, prepare=prepare)
    _, xlsx, output, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "Material", "Timber")
    apply(paths)
    again = tmp_path / "again.ifc"
    result = update_ifc_from_excel(str(output), str(xlsx), str(again))
    f = ifcopenshell.open(str(again))
    assert util.get_material(f.by_guid(b)).Name == "Timber"
    assert result["changes"]["applied"] == 0


def test_split_identity_history_survives_repeated_imports(tmp_path):
    paths = export(tmp_path)
    _, xlsx, output, (a, b, typ, _) = paths
    edit(xlsx, "Elements", a, "Type Name", "A type")
    edit(xlsx, "Elements", b, "Type Name", "B type")
    f, _ = apply(paths, split_shared_types=True)
    expected = [util.get_type(f.by_guid(g)).GlobalId for g in (a, b)]
    previous = output
    for index in range(2):
        following = tmp_path / f"again-{index}.ifc"
        result = update_ifc_from_excel(str(previous), str(xlsx), str(following), split_shared_types=True)
        f = ifcopenshell.open(str(following))
        assert [util.get_type(f.by_guid(g)).GlobalId for g in (a, b)] == expected
        assert result["changes"]["created_types"] == []
        assert result["changes"]["applied"] == 0
        previous = following
