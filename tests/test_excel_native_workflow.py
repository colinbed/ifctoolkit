"""Actual-target round trips, including schema validation and association isolation."""
import hashlib
import ifcopenshell
import ifcopenshell.api
import ifcopenshell.util.element as util
import pytest
from ifcopenshell.guid import new as guid
from openpyxl import load_workbook
from app import extract_to_excel, update_ifc_from_excel, ExcelWritebackValidationError
from ifc_excel_schema import schema_reference
from ifc_excel_targets import FLOOR_REFERENCE, FLOOR_NAME
from test_excel_name_resolution import exported, edit


def run(paths, **kwargs):
    source, workbook, output, *_ = paths
    result = update_ifc_from_excel(str(source), str(workbook), str(output), **kwargs)
    return ifcopenshell.open(str(output)), result


def test_native_names_descriptions_and_protocol(exported):
    _, workbook, _, guids, type_guid = exported
    edit(workbook, [("Elements", guids[0], "Instance Name", "Instance"),
                    ("Elements", guids[0], "Instance Description", "Instance description"),
                    ("Elements", guids[0], "Type Name", "Type"),
                    ("Types", type_guid, "Type Name", "Type"),
                    ("Types", type_guid, "Type Description", "Type description")])
    model, result = run(exported)
    assert model.by_guid(guids[0]).Name == "Instance"
    assert model.by_guid(guids[0]).Description == "Instance description"
    assert model.by_guid(guids[1]).Description is None
    assert util.get_type(model.by_guid(guids[0])).Description == "Type description"
    assert model.by_guid(type_guid).Name == "Type"
    assert result['changes']['applied'] == 4


def test_types_elements_conflict(exported):
    _, workbook, output, guids, type_guid = exported
    edit(workbook, [("Elements", guids[0], "Type Description", "A"), ("Types", type_guid, "Type Description", "B")])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        run(exported)
    assert caught.value.findings['errors'][0]['target_guid'] == type_guid
    assert not output.exists()


def test_simple_shared_material_isolation(exported):
    _, workbook, _, guids, _ = exported
    edit(workbook, [("Elements", guids[0], "Material", "Aluminium")])
    model, result = run(exported)
    assert util.get_material(model.by_guid(guids[0])).Name == "Aluminium"
    assert util.get_material(model.by_guid(guids[1])).Name == "Steel"
    assert result['changes']['applied'] == 1
    assert not util.get_psets(model.by_guid(guids[0])).get('Material')


def test_type_material_conflicts_and_scope(exported):
    source, workbook, _, guids, type_guid = exported
    model = ifcopenshell.open(str(source))
    material = util.get_material(model.by_guid(guids[0]))
    for rel in list(model.by_type('IfcRelAssociatesMaterial')):
        model.remove(rel)
    model.create_entity('IfcRelAssociatesMaterial', GlobalId=guid(), RelatedObjects=[model.by_guid(type_guid)], RelatingMaterial=material)
    unrelated = model.create_entity('IfcBeam', GlobalId=guid())
    unrelated_guid = unrelated.GlobalId
    model.create_entity('IfcRelAssociatesMaterial', GlobalId=guid(), RelatedObjects=[unrelated], RelatingMaterial=material)
    model.write(str(source))
    extract_to_excel(str(source), str(workbook))
    edit(workbook, [('Elements', guids[0], 'Material', 'A'), ('Elements', guids[1], 'Material', 'B')])
    with pytest.raises(ExcelWritebackValidationError):
        run(exported)
    edit(workbook, [('Elements', guids[1], 'Material', 'A')])
    model, _ = run(exported)
    assert all(util.get_material(model.by_guid(g)).Name == 'A' for g in guids)
    assert util.get_material(model.by_guid(unrelated_guid)).Name == 'Steel'


def test_structured_material_preserved_and_edit_rejected(exported):
    source, workbook, _, guids, _ = exported
    model = ifcopenshell.open(str(source))
    for rel in list(model.by_type('IfcRelAssociatesMaterial')):
        model.remove(rel)
    steel = model.create_entity('IfcMaterial', Name='Steel')
    layer = model.create_entity('IfcMaterialLayer', Material=steel, LayerThickness=0.2)
    layers = model.create_entity('IfcMaterialLayerSet', MaterialLayers=[layer], LayerSetName='Assembly')
    model.create_entity('IfcRelAssociatesMaterial', GlobalId=guid(), RelatedObjects=[model.by_guid(guids[0])], RelatingMaterial=layers)
    model.write(str(source))
    extract_to_excel(str(source), str(workbook))
    revised, _ = run(exported)
    assert util.get_material(revised.by_guid(guids[0])).is_a('IfcMaterialLayerSet')
    edit(workbook, [('Elements', guids[0], 'Material', 'Flattened')])
    with pytest.raises(ExcelWritebackValidationError, match='validation'):
        run(exported)


def test_storey_attributes_properties_and_classification(tmp_path):
    model = ifcopenshell.file(schema='IFC4')
    project = model.create_entity('IfcProject', GlobalId=guid(), Name='P')
    building = model.create_entity('IfcBuilding', GlobalId=guid(), Name='B')
    storey = model.create_entity('IfcBuildingStorey', GlobalId=guid(), Name='L0', Elevation=3.0)
    component = model.create_entity('IfcBeam', GlobalId=guid(), Name='Beam')
    for parent, child in [(project, building), (building, storey)]:
        model.create_entity('IfcRelAggregates', GlobalId=guid(), RelatingObject=parent, RelatedObjects=[child])
    model.create_entity('IfcRelContainedInSpatialStructure', GlobalId=guid(), RelatingStructure=storey, RelatedElements=[component])
    # Preserve an existing property representation separately from classification.
    pset = ifcopenshell.api.run('pset.add_pset', model, product=storey, name='COBie_Floor')
    ifcopenshell.api.run('pset.edit_pset', model, pset=pset, properties={'Category': 'Original category'})
    sg, cg, bg = storey.GlobalId, component.GlobalId, building.GlobalId
    paths = [tmp_path / n for n in ['source.ifc','book.xlsx','output.ifc']]
    model.write(str(paths[0])); digest = hashlib.sha256(paths[0].read_bytes()).hexdigest()
    extract_to_excel(str(paths[0]), str(paths[1]))
    edit(paths[1], [('BuildingStoreys', sg, field, value) for field, value in [
        ('Name','First'), ('Description','Description'), ('Long Name','Long'),
        ('COBie_Floor.Category','Category'), ('COBie_FloorAttributes.HeightToHaunch', '4.2'),
        (FLOOR_REFERENCE, 'EF_30'), (FLOOR_NAME,'Floor category')]])
    revised, _ = run(paths)
    s = revised.by_guid(sg)
    assert (s.Name, s.Description, s.LongName, s.Elevation) == ('First','Description','Long',3.)
    assert util.get_aggregate(s).GlobalId == bg
    assert util.get_container(revised.by_guid(cg)).GlobalId == sg
    assert revised.by_guid(cg).Name == 'Beam'
    assert util.get_psets(s)['COBie_Floor']['Category'] == 'Category'
    assert util.get_psets(s)['COBie_FloorAttributes']['HeightToHaunch'] == '4.2'
    ref = next(r.RelatingClassification for r in s.HasAssociations if r.is_a('IfcRelAssociatesClassification'))
    assert (ref.Identification, ref.Name, ref.ReferencedSource.Name) == ('EF_30','Floor category','COBie Floors')
    assert hashlib.sha256(paths[0].read_bytes()).hexdigest() == digest


@pytest.mark.parametrize('schema', ['IFC2X3', 'IFC4', 'IFC4X3_ADD2'])
def test_schema_lists_and_dependent_validations(tmp_path, schema):
    model = ifcopenshell.file(schema=schema)
    model.create_entity('IfcProject', GlobalId=guid(), Name='P')
    beam = model.create_entity('IfcBeam', GlobalId=guid())
    typ = model.create_entity('IfcBeamType', GlobalId=guid(), PredefinedType='BEAM')
    model.create_entity('IfcRelDefinesByType', GlobalId=guid(), RelatedObjects=[beam], RelatingType=typ)
    paths = [tmp_path / n for n in ['source.ifc','book.xlsx','output.ifc']]
    model.write(str(paths[0])); extract_to_excel(str(paths[0]), str(paths[1]))
    wb = load_workbook(paths[1]); refs = schema_reference(schema)
    assert refs['IfcElement']['abstract'] and not refs['IfcElement']['selectable']
    assert refs['IfcWall']['selectable']  # absent from source
    assert refs['IfcBeamType']['enum'] != refs['IfcWallType']['enum']
    for sheet in ['SchemaReference', 'SchemaChoices']:
        assert wb[sheet].sheet_state == 'visible'
    for sheet in ['Elements', 'Types']:
        dvs = list(wb[sheet].data_validations.dataValidation)
        assert len(dvs) == 2 and all(dv.sqref for dv in dvs)
        assert any('INDIRECT' in dv.formula1 for dv in dvs)
    assert wb.defined_names['IFC_Enum_IfcBeamType'].attr_text.startswith("'SchemaChoices'!$D$")
    assert 'IFC_Types' in wb.defined_names and 'IFC_Occurrences' in wb.defined_names
    wb.save(paths[1]); wb.close()
    reopened = load_workbook(paths[1]); assert len(reopened['Types'].data_validations.dataValidation) == 2; reopened.close()
    if schema == 'IFC2X3':
        wb = load_workbook(paths[1]); cols = {c.value:c.column for c in wb['Elements'][1]}
        assert wb['Elements'].cell(2,cols['Predefined Type']).value == 'Unavailable'; wb.close()
    edit(paths[1], [('Types', typ.GlobalId, 'Predefined Type', 'INVALID_PASTED')])
    with pytest.raises(ExcelWritebackValidationError):
        run(paths)
    assert not paths[2].exists()


def test_type_conversion_preserves_geometry_relationships(exported):
    source, workbook, _, guids, type_guid = exported
    before = ifcopenshell.open(str(source)); placement = before.by_guid(guids[0]).ObjectPlacement.id(); representation = before.by_guid(guids[0]).Representation.id()
    edit(workbook, [('Types', type_guid, 'IFC Entity', 'IfcWallType'), ('Types', type_guid, 'Predefined Type', 'NOTDEFINED')])
    model, _ = run(exported)
    assert model.by_guid(type_guid).is_a() == 'IfcWallType'
    assert all(model.by_guid(g).is_a('IfcWall') for g in guids)
    assert all(util.get_type(model.by_guid(g)).GlobalId == type_guid for g in guids)
    assert model.by_guid(guids[0]).ObjectPlacement.id() == placement
    assert model.by_guid(guids[0]).Representation.id() == representation


def test_legacy_headers_and_conflicting_aliases(exported):
    _, workbook, _, guids, type_guid = exported
    wb = load_workbook(workbook); ws = wb['Elements']; cols = {c.value:c.column for c in ws[1]}
    for old, new in [('Instance Name','OccurrenceName'),('Type Name','TypeName'),('Instance Description','TypeDescription')]:
        ws.cell(1,cols[old],new)
    wb.save(workbook); wb.close()
    edit(workbook, [('Elements',guids[0],'OccurrenceName','Legacy'), ('Elements',guids[0],'TypeName','Legacy type'), ('Elements',guids[0],'TypeDescription','Legacy occurrence description')])
    model, _ = run(exported)
    assert model.by_guid(guids[0]).Name == 'Legacy'
    assert model.by_guid(guids[0]).Description == 'Legacy occurrence description'
    assert model.by_guid(type_guid).Name == 'Legacy type'
    # Two explicit aliases for one canonical target must conflict.
    wb = load_workbook(workbook); ws=wb['Elements']; col=ws.max_column+1; ws.cell(1,col,'Instance Name'); ws.cell(2,col,'Conflict'); wb.save(workbook); wb.close()
    with pytest.raises(ExcelWritebackValidationError):
        run(exported)


def test_userdefined_requires_companion(exported):
    _, workbook, _, _, type_guid = exported
    edit(workbook, [('Types', type_guid, 'Predefined Type', 'USERDEFINED')])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        run(exported)
    assert any('ElementType' in e['message'] for e in caught.value.findings['errors'])
    edit(workbook, [('Types', type_guid, 'Element Type', 'Custom type')])
    model, _ = run(exported)
    assert model.by_guid(type_guid).PredefinedType == 'USERDEFINED'
    assert model.by_guid(type_guid).ElementType == 'Custom type'


def test_shared_class_conflicts_and_abstract_targets(exported):
    _, workbook, _, guids, type_guid = exported
    edit(workbook, [('Elements', guids[0], 'IFC Entity', 'IfcWall'),
                    ('Types', type_guid, 'IFC Entity', 'IfcSlabType')])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        run(exported)
    assert caught.value.findings['changes']['conflicting'] > 0
    edit(workbook, [('Types', type_guid, 'IFC Entity', 'IfcElementType')])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        run(exported)
    assert any('Unsupported conversion' in e['message'] for e in caught.value.findings['errors'])


@pytest.mark.parametrize('schema', ['IFC2X3', 'IFC4'])
def test_floor_shared_properties_and_classification_are_isolated(tmp_path, schema):
    model = ifcopenshell.file(schema=schema)
    model.create_entity('IfcProject', GlobalId=guid(), Name='P')
    floors = [model.create_entity('IfcBuildingStorey', GlobalId=guid(), Name='Level') for _ in range(2)]
    prop = model.create_entity('IfcPropertySingleValue', Name='Category', NominalValue=model.create_entity('IfcLabel','Original'))
    pset = model.create_entity('IfcPropertySet', GlobalId=guid(), Name='COBie_Floor', HasProperties=[prop])
    model.create_entity('IfcRelDefinesByProperties', GlobalId=guid(), RelatedObjects=floors, RelatingPropertyDefinition=pset)
    system = model.create_entity('IfcClassification', Name='COBie Floors')
    ref = model.create_entity('IfcClassificationReference', Name='Original name', ReferencedSource=system,
                              **({'ItemReference':'Original'} if schema == 'IFC2X3' else {'Identification':'Original'}))
    model.create_entity('IfcRelAssociatesClassification', GlobalId=guid(), RelatedObjects=floors, RelatingClassification=ref)
    floor_guids = [f.GlobalId for f in floors]
    paths = [tmp_path / n for n in ['source.ifc','book.xlsx','output.ifc']]
    model.write(str(paths[0])); extract_to_excel(str(paths[0]),str(paths[1]))
    edit(paths[1], [('BuildingStoreys',floor_guids[0],'COBie_Floor.Category','Updated'),
                    ('BuildingStoreys',floor_guids[0],FLOOR_REFERENCE,'Updated')])
    revised, _ = run(paths)
    from ifc_excel_targets import read_association
    first, second = [revised.by_guid(g) for g in floor_guids]
    assert util.get_psets(first)['COBie_Floor']['Category'] == 'Updated'
    assert util.get_psets(second)['COBie_Floor']['Category'] == 'Original'
    assert read_association(first, dict(kind='floor_classification',attribute='Reference')) == 'Updated'
    assert read_association(second, dict(kind='floor_classification',attribute='Reference')) == 'Original'
