import hashlib

import ifcopenshell
import ifcopenshell.api
import ifcopenshell.util.element
import pytest
from ifcopenshell.guid import new as new_guid
from openpyxl import load_workbook

from app import ExcelWritebackValidationError, extract_to_excel, update_ifc_from_excel
from ifc_protocol import default_protocol_config


@pytest.fixture
def exported(tmp_path):
    model = ifcopenshell.file(schema="IFC4")
    model.create_entity("IfcProject", GlobalId=new_guid(), Name="Project", LongName="P1")
    shared = model.create_entity("IfcBeamType", GlobalId=new_guid(), Name="PB-2", PredefinedType="BEAM")
    beams = [model.create_entity("IfcBeam", GlobalId=new_guid(), Name="BEAM", ObjectType="Beam") for _ in range(2)]
    ifcopenshell.api.run("type.assign_type", model, related_objects=beams, relating_type=shared)
    material = ifcopenshell.api.run("material.add_material", model, name="Steel")
    ifcopenshell.api.run("material.assign_material", model, products=beams, material=material)
    # A real representation and placement must survive name edits unchanged.
    point = model.create_entity("IfcCartesianPoint", Coordinates=(0., 0., 0.))
    placement = model.create_entity("IfcAxis2Placement3D", Location=point)
    beams[0].ObjectPlacement = model.create_entity("IfcLocalPlacement", RelativePlacement=placement)
    context = ifcopenshell.api.run("context.add_context", model, context_type="Model")
    solid = model.create_entity("IfcBoundingBox", Corner=point, XDim=1., YDim=2., ZDim=3.)
    shape = model.create_entity("IfcShapeRepresentation", ContextOfItems=context,
                                RepresentationIdentifier="Box", RepresentationType="BoundingBox", Items=[solid])
    beams[0].Representation = model.create_entity("IfcProductDefinitionShape", Representations=[shape])
    pset = ifcopenshell.api.run("pset.add_pset", model, product=beams[0], name="COBie_Component")
    ifcopenshell.api.run("pset.edit_pset", model, pset=pset, properties={"Name": "COBie name"})
    source, workbook, output = (tmp_path / name for name in ("source.ifc", "source.xlsx", "output.ifc"))
    model.write(str(source))
    extract_to_excel(str(source), str(workbook), plan_payload={"protocol": default_protocol_config()})
    return source, workbook, output, [b.GlobalId for b in beams], shared.GlobalId


def edit(path, changes):
    wb = load_workbook(path)
    for sheet_name, guid, field, value in changes:
        sheet = wb[sheet_name]
        columns = {c.value: c.column for c in sheet[1]}
        row = next(i for i in range(2, sheet.max_row + 1) if sheet.cell(i, columns.get("GlobalId", columns.get("Type GlobalId"))).value == guid)
        sheet.cell(row, columns[field]).value = value
    wb.save(path)
    wb.close()


@pytest.mark.parametrize("sheets", [("Elements",), ("ProtocolData",), ("Elements", "ProtocolData")])
def test_name_edits_resolve_and_survive_reopen(exported, sheets):
    source, workbook, output, guids, type_guid = exported
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    changes = []
    for sheet in sheets:
        changes.append((sheet, guids[0], "Instance Name" if sheet == "Elements" else "Asset Name", "EST-BEM-00001"))
        for guid in guids:
            changes.append((sheet, guid, "Type Name" if sheet == "Elements" else "Type Name", "EST_S_Beam_Beam001"))
    changes.append(("COBieMapping", guids[0], "COBie_Component.Name", "Edited COBie name"))
    edit(workbook, changes)
    result = update_ifc_from_excel(str(source), str(workbook), str(output))
    before, after = ifcopenshell.open(str(source)), ifcopenshell.open(str(output))
    assert after.by_guid(guids[0]).Name == "EST-BEM-00001"
    assert after.by_guid(guids[1]).Name == "BEAM"
    assert after.by_guid(type_guid).Name == "EST_S_Beam_Beam001"
    assert result["changes"]["applied"] == 3  # occurrence, shared type, property
    assert result["changes"]["conflicting"] == 0
    assert result["changes"]["unchanged"] > 0
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    # Compare every original entity by STEP id as well as product GlobalId.
    for old in before:
        new = after.by_id(old.id())
        for key, value in old.get_info().items():
            if key == "Name" and getattr(old, "GlobalId", None) in {guids[0], type_guid}:
                continue
            if key == "NominalValue" and old.is_a("IfcPropertySingleValue") and old.Name == "Name":
                continue
            def stable(v):
                if isinstance(v, ifcopenshell.entity_instance):
                    return (v.is_a(), v.id()) if v.id() else str(v)
                if isinstance(v, tuple):
                    return tuple(stable(item) for item in v)
                return v
            assert stable(new.get_info()[key]) == stable(value)
    for guid in guids:
        assert ifcopenshell.util.element.get_type(after.by_guid(guid)).GlobalId == type_guid
    assert ifcopenshell.util.element.get_psets(after.by_guid(guids[0]))["COBie_Component"]["Name"] == "Edited COBie name"


def test_cross_sheet_conflict_is_structured_and_atomic(exported):
    source, workbook, output, guids, _ = exported
    edit(workbook, [("Elements", guids[0], "Instance Name", "ONE"),
                    ("ProtocolData", guids[0], "Asset Name", "TWO")])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        update_ifc_from_excel(str(source), str(workbook), str(output))
    errors = caught.value.findings["errors"]
    assert len(errors) == 1
    groups = errors[0]["requested_values"]
    affected = [row for group in groups for row in group["affected"]]
    assert {row["sheet"] for row in affected} == {"Elements", "ProtocolData"}
    assert {group["value"] for group in groups} == {"ONE", "TWO"}
    assert all(row["row"] >= 2 and row["global_id"] == guids[0] for row in affected)
    assert not output.exists()


@pytest.mark.parametrize("sheet,column", [("Elements", "Type Name"), ("ProtocolData", "Type Name")])
def test_shared_type_conflict_reports_type_and_occurrence_rows(exported, sheet, column):
    source, workbook, output, guids, type_guid = exported
    edit(workbook, [(sheet, guid, column, f"TYPE-{i}") for i, guid in enumerate(guids)])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        update_ifc_from_excel(str(source), str(workbook), str(output))
    errors = caught.value.findings["errors"]
    assert len(errors) == 1
    affected = [row for group in errors[0]["requested_values"] for row in group["affected"]]
    assert {row["global_id"] for row in affected} == set(guids)
    assert errors[0]["target_guid"] == type_guid
    assert len({row["row"] for row in affected}) == 2
    assert type_guid in errors[0]["message"]
    assert len(errors[0]["message"]) < 200
    assert not output.exists()


@pytest.mark.parametrize("blank", [None, "   "])
def test_blank_names_preserve_source(exported, blank):
    source, workbook, output, guids, type_guid = exported
    edit(workbook, [(sheet, guid, column, blank) for guid in guids
                    for sheet, column in [("Elements", "Instance Name"), ("Elements", "Type Name"),
                                           ("ProtocolData", "Asset Name"), ("ProtocolData", "Type Name")]])
    result = update_ifc_from_excel(str(source), str(workbook), str(output))
    after = ifcopenshell.open(str(output))
    assert all(after.by_guid(g).Name == "BEAM" for g in guids)
    assert after.by_guid(type_guid).Name == "PB-2"
    assert result["changes"]["applied"] == 0


def test_one_occurrence_can_rename_shared_type(exported):
    source, workbook, output, guids, type_guid = exported
    edit(workbook, [("Elements", guids[0], "Type Name", "SHARED")])
    result = update_ifc_from_excel(str(source), str(workbook), str(output))
    model = ifcopenshell.open(str(output))
    assert all(ifcopenshell.util.element.get_type(model.by_guid(g)).Name == "SHARED" for g in guids)
    assert model.by_guid(type_guid).Name == "SHARED"
    assert result["changes"]["applied"] == 1


def test_shared_type_cross_sheet_conflict(exported):
    source, workbook, output, guids, type_guid = exported
    edit(workbook, [("Elements", guids[0], "Type Name", "ONE"),
                    ("ProtocolData", guids[1], "Type Name", "TWO")])
    with pytest.raises(ExcelWritebackValidationError) as caught:
        update_ifc_from_excel(str(source), str(workbook), str(output))
    assert {e["target_guid"] for e in caught.value.findings["errors"]} == {type_guid}
    assert caught.value.findings["changes"]["applied"] == 0
    assert caught.value.findings["changes"]["conflicting"] == 1
    assert not output.exists()


def test_legacy_object_type_edit_with_unchanged_protocol(exported):
    source, workbook, output, guids, _ = exported
    edit(workbook, [("Elements", guids[0], "Object Type", "Edited ObjectType")])
    update_ifc_from_excel(str(source), str(workbook), str(output))
    assert ifcopenshell.open(str(output)).by_guid(guids[0]).ObjectType == "Edited ObjectType"


def test_types_name_edit_is_independent_of_apply_change(exported):
    source, workbook, output, _, type_guid = exported
    edit(workbook, [("Types", type_guid, "Type Name", "TYPE EDIT")])
    update_ifc_from_excel(str(source), str(workbook), str(output))
    assert ifcopenshell.open(str(output)).by_guid(type_guid).Name == "TYPE EDIT"


@pytest.mark.parametrize("legacy,protocol,conflict", [
    ("Edited", None, False), (None, "Edited", False),
    ("Edited", "Edited", False), ("ONE", "TWO", True),
])
def test_overlapping_property_targets(exported, legacy, protocol, conflict):
    source, workbook, output, guids, _ = exported
    config = default_protocol_config()
    config["entities"][0]["fields"].append({
        "id": "cobie_name", "label": "COBie Name", "editable": True,
        "source": {"kind": "property", "pset": "COBie_Component", "property": "Name"},
        "write": {"enabled": True, "target": {"kind": "occurrence_property", "pset": "COBie_Component", "property": "Name"}},
    })
    extract_to_excel(str(source), str(workbook), plan_payload={"protocol": config})
    changes = []
    if legacy:
        changes.append(("COBieMapping", guids[0], "COBie_Component.Name", legacy))
    if protocol:
        changes.append(("ProtocolData", guids[0], "COBie Name", protocol))
    edit(workbook, changes)
    if conflict:
        with pytest.raises(ExcelWritebackValidationError):
            update_ifc_from_excel(str(source), str(workbook), str(output))
        assert not output.exists()
    else:
        update_ifc_from_excel(str(source), str(workbook), str(output))
        model = ifcopenshell.open(str(output))
        assert model.by_guid(guids[0]).Name == "BEAM"
        assert ifcopenshell.util.element.get_psets(model.by_guid(guids[0]))["COBie_Component"]["Name"] == "Edited"


def test_cobie_type_name_does_not_rename_native_type(exported):
    source, workbook, output, guids, type_guid = exported
    config = default_protocol_config()
    config["entities"][0]["fields"].append({
        "id": "cobie_type_name", "label": "COBie Type Name", "editable": True,
        "source": {"kind": "type_property", "pset": "COBie_Type", "property": "Name"},
        "write": {"enabled": True, "create_if_missing": True},
    })
    extract_to_excel(str(source), str(workbook), plan_payload={"protocol": config})
    edit(workbook, [("ProtocolData", guids[0], "COBie Type Name", "COBIE-TYPE")])
    update_ifc_from_excel(str(source), str(workbook), str(output))
    model = ifcopenshell.open(str(output))
    assert model.by_guid(type_guid).Name == "PB-2"
    assert ifcopenshell.util.element.get_psets(model.by_guid(type_guid))["COBie_Type"]["Name"] == "COBIE-TYPE"
