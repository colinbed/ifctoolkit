"""Executable Excel choices from the installed EXPRESS schema, never a static enum union."""
from functools import lru_cache
import ifcopenshell
import ifcopenshell.ifcopenshell_wrapper as wrapper
import ifcopenshell.util.type
import ifcopenshell.util.doc

REFERENCE_VERSION = "2"


def is_subtype(declaration, name):
    while declaration:
        if declaration.name() == name:
            return True
        declaration = declaration.supertype()
    return False


@lru_cache(maxsize=8)
def schema_reference(schema):
    definition = wrapper.schema_by_name(schema)
    rows = {}
    for decl in definition.entities():
        category = ("type" if is_subtype(decl, "IfcTypeObject") else
                    "occurrence" if is_subtype(decl, "IfcProduct") else "reference")
        attributes = {a.name(): a for a in decl.all_attributes()}
        enum = ()
        if "PredefinedType" in attributes:
            enum = tuple(attributes["PredefinedType"].type_of_attribute().declared_type().enumeration_items())
        corresponding = ()
        try:
            corresponding = tuple(ifcopenshell.util.type.get_applicable_entities(decl.name(), schema) if category == "type"
                                  else ifcopenshell.util.type.get_applicable_types(decl.name(), schema))
        except (RuntimeError, KeyError):
            pass
        # Restrict conversions to element families supported by root.reassign_class.
        selectable = (is_subtype(decl, "IfcElement") or is_subtype(decl, "IfcElementType")) and not decl.is_abstract()
        base = "IfcElementType" if category == "type" else "IfcElement"
        common = {a.name() for a in definition.declaration_by_name(base).all_attributes()}
        editable = {"PredefinedType", "ElementType", "ObjectType"}
        if any(not a.optional() and a.name() not in common | editable for a in attributes.values()):
            selectable = False
        if schema.upper() == "IFC2X3":
            url = f"https://standards.buildingsmart.org/IFC/RELEASE/IFC2x3/TC1/HTML/ifcsharedbldgelements/lexical/{decl.name().lower()}.htm"
        else:
            release = "IFC4x3_ADD2" if schema.upper().startswith("IFC4X3") else "IFC4/ADD2_TC1"
            url = f"https://standards.buildingsmart.org/IFC/RELEASE/{release}/HTML/lexical/{decl.name().lower()}.htm"
        doc_schema = "IFC4X3" if schema.upper().startswith("IFC4X3") else schema.upper()
        try:
            doc = ifcopenshell.util.doc.get_entity_doc(doc_schema, decl.name(), recursive=False) or {}
        except (KeyError, AssertionError):
            doc = {}
        url = doc.get("spec_url") or url
        rows[decl.name()] = dict(entity=decl.name(), category=category, abstract=decl.is_abstract(),
                                 selectable=selectable, attributes=tuple(attributes), enum=enum,
                                 corresponding=corresponding, url=url, description=doc.get("description", "Documentation unavailable in installed IfcOpenShell"))
    return rows


def validate_native(schema, entity, attribute, value):
    info = schema_reference(schema).get(entity)
    if not info or attribute not in info["attributes"]:
        return f"{entity} has no native {attribute} in {schema}."
    if attribute == "PredefinedType" and value not in info["enum"]:
        return f"{value!r} is not a valid {entity}.PredefinedType in {schema}; choose from the schema reference."
    return None


def add_reference_sheets(workbook, sheet, schema):
    from openpyxl.workbook.defined_name import DefinedName
    refs = schema_reference(schema)
    ws = sheet("SchemaReference", ["Schema", "IFC Entity", "Category", "Abstract", "Selectable", "Correspondence", "Native PredefinedType", "Enumeration", "Documentation URL", "Reference Version", "Description"])
    for info in refs.values():
        if info["category"] == "reference":
            continue
        ws.append([schema, info["entity"], info["category"], info["abstract"], info["selectable"],
                   ", ".join(info["corresponding"]), "PredefinedType" in info["attributes"],
                   ", ".join(info["enum"]), info["url"], REFERENCE_VERSION, info["description"]])
    values = sheet("SchemaChoices", ["Occurrence Entities", "Type Entities", "Entity", "Predefined Type"])
    occurrences = [n for n, i in refs.items() if i["selectable"] and i["category"] == "occurrence"]
    types = [n for n, i in refs.items() if i["selectable"] and i["category"] == "type"]
    enums = [(n, v) for n, i in refs.items() for v in (i["enum"] or ("",))]
    for index in range(max(len(occurrences), len(types), len(enums))):
        values.append([occurrences[index] if index < len(occurrences) else None,
                       types[index] if index < len(types) else None,
                       *(enums[index] if index < len(enums) else (None, None))])
    def define(name, column, start, end):
        item = DefinedName(name, attr_text=f"'SchemaChoices'!${column}${start}:${column}${end}")
        if hasattr(workbook.defined_names, "add"):
            workbook.defined_names.add(item)
        else:
            workbook.defined_names.append(item)
    define("IFC_Occurrences", "A", 2, len(occurrences) + 1)
    define("IFC_Types", "B", 2, len(types) + 1)
    offset = 2
    for name, info in refs.items():
        length = max(1, len(info["enum"]))
        define("IFC_Enum_" + name, "D", offset, offset + length - 1)
        offset += length


def add_native_validations(ws, columns, count, category):
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    if not count:
        return
    entity = get_column_letter(columns.index("IFC Entity") + 1)
    predefined = get_column_letter(columns.index("Predefined Type") + 1)
    for letter, formula in ((entity, "=IFC_Types" if category == "type" else "=IFC_Occurrences"),
                            (predefined, f'=INDIRECT("IFC_Enum_"&${entity}2)')):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.errorTitle = "Invalid IFC schema value"
        dv.error = "Choose a value from the schema reference for this IFC entity."
        dv.showErrorMessage = True
        dv.errorStyle = "stop"
        dv.add(f"{letter}2:{letter}{count + 1}")
        ws.data_validations.append(dv)


def read_class_overrides(path, model):
    """Retain explicit class cells when re-extracting into an existing workbook."""
    from pathlib import Path
    from openpyxl import load_workbook
    result = {}
    if not Path(path).exists():
        return result
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in ("Elements", "Types"):
            if sheet not in wb:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            columns = next(rows, ())
            for values in rows:
                row = dict(zip(columns, values))
                guid = row.get("Type GlobalId") if sheet == "Types" and "Type GlobalId" in row else row.get("GlobalId")
                try:
                    obj = model.by_guid(str(guid))
                except RuntimeError:
                    continue
                requests = {}
                for header, attr in (("IFC Entity", None), ("Predefined Type", "PredefinedType")):
                    value = row.get(header)
                    old = obj.is_a() if attr is None else getattr(obj, attr, None)
                    if value and value != "Unavailable" and value != old:
                        requests[header] = value
                if str(row.get("ApplyChange") or "").lower() in {"yes", "y", "1", "true"}:
                    for legacy, native in (("TargetEntity", "IFC Entity"), ("TargetPredefinedType", "Predefined Type")):
                        if row.get(legacy) and row[legacy] != "UNCHANGED":
                            requests[native] = row[legacy]
                if requests:
                    result[(sheet, guid)] = requests
    finally:
        wb.close()
    return result
